"""Extract published Nyaungdon annual river-stage context for provenance only.

This extracts the first historical Panhlaing River table in the openly published
PLOS S1 workbook. The series is historical Maubin-District context, not a Maubin
Township live gauge, tide series, model feature, forecast, alert, or prediction.
"""

from __future__ import annotations

import json
from datetime import UTC, date, datetime
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path("/home/ubuntu/nationwide-data/maubin_local_water/plos_nyaungdon_s1_hydro_data.xlsx")
OUTPUT = Path("/home/ubuntu/deltawatch-model-outputs/nyaungdon_panhlaing_historical_stage_context.json")


def parse_date(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value.strip(), fmt).date().isoformat()
            except ValueError:
                continue
    return None


def number(value: object) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return None


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = workbook["Observed_data"]
    records: list[dict] = []
    for row in sheet.iter_rows(min_row=3, max_col=3, values_only=True):
        observed_date = parse_date(row[0])
        high_stage_ft = number(row[1])
        danger_stage_ft = number(row[2])
        if observed_date and high_stage_ft is not None and danger_stage_ft is not None:
            records.append(
                {
                    "observed_date": observed_date,
                    "annual_high_stage_ft": high_stage_ft,
                    "published_danger_stage_ft": danger_stage_ft,
                    "above_or_equal_published_danger_stage": high_stage_ft >= danger_stage_ft,
                }
            )
    if not records:
        raise RuntimeError("No structured Panhlaing annual-stage records found in published workbook")
    records.sort(key=lambda record: record["observed_date"])
    result = {
        "schema": "deltawatch-nyaungdon-panhlaing-historical-stage-context-v1",
        "generated_at_utc": datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "source": {
            "citation": "Khaing et al. (2019), PLOS ONE 14(11): e0224558, S1 Dataset",
            "doi": "https://doi.org/10.1371/journal.pone.0224558",
            "workbook_url": "https://journals.plos.org/plosone/article/file?type=supplementary&id=10.1371/journal.pone.0224558.s001",
            "workbook_sha256": "ce1c7875ae66b52e5461a26436528fccd794d5037e4b808192862122376a7dc6",
            "sheet": "Observed_data",
            "published_table_label": "Panhling_River; Highest Wl (ft); Danger level",
        },
        "geographic_scope": {
            "station_label": "Nyaungdon station on the Panhlaing River",
            "administrative_context": "Nyaungdon Township, Maubin District, Ayeyarwady Region",
            "relationship_to_maubin_township": "Nearby district-level historical context only; not a Maubin Township gauge equivalence.",
        },
        "coverage": {
            "record_count": len(records),
            "first_observed_date": records[0]["observed_date"],
            "last_observed_date": records[-1]["observed_date"],
            "danger_or_higher_record_count": sum(record["above_or_equal_published_danger_stage"] for record in records),
        },
        "records": records,
        "operational_authorization": {
            "historical_context_allowed": True,
            "live_river_stage_allowed": False,
            "tide_or_coastal_water_allowed": False,
            "model_feature_allowed": False,
            "reason": "The source is an annual historical study table with no current telemetry, live-feed endpoint, verified common datum, or issue-time/prospective lineage.",
        },
        "safety": "Monitoring only: historical provenance context only; no model fit, score, probability, prediction, forecast, alert, or life-safety decision.",
    }
    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT), "coverage": result["coverage"], "model_feature_allowed": False}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
