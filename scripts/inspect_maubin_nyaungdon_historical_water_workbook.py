"""Inspect published Nyaungdon historical water evidence without modelling.

The PLOS supplementary workbook contains historical study inputs. This utility
records sheet layouts and non-empty cell coverage for provenance review only. It
does not construct operational features, scores, probabilities, predictions,
forecasts, alerts, or life-safety decisions.
"""

from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path("/home/ubuntu/nationwide-data/maubin_local_water/plos_nyaungdon_s1_hydro_data.xlsx")
OUTPUT = Path("/home/ubuntu/deltawatch-model-outputs/maubin_nyaungdon_historical_water_workbook_audit.json")


def display(value: object) -> str | None:
    if value is None:
        return None
    return str(value).strip() or None


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheets: list[dict] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        nonempty = sum(1 for row in rows for value in row if value is not None)
        sample_rows = [[display(value) for value in row[:12]] for row in rows[:8]]
        sheets.append(
            {
                "sheet_name": sheet.title,
                "max_rows": sheet.max_row,
                "max_columns": sheet.max_column,
                "nonempty_cell_count": nonempty,
                "sample_rows": sample_rows,
            }
        )
    result = {
        "schema": "deltawatch-maubin-nyaungdon-historical-water-workbook-audit-v1",
        "generated_at_utc": datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "workbook_path": str(WORKBOOK),
        "workbook_bytes": WORKBOOK.stat().st_size,
        "sha256": "ce1c7875ae66b52e5461a26436528fccd794d5037e4b808192862122376a7dc6",
        "citation": "Khaing et al. (2019), PLOS ONE 14(11): e0224558, S1 Dataset",
        "source_url": "https://journals.plos.org/plosone/article/file?type=supplementary&id=10.1371/journal.pone.0224558.s001",
        "geographic_scope": "Nyaungdon Township, Maubin District, Ayeyarwady Region; not automatically interchangeable with a Maubin Township gauge.",
        "evidence_class": "published historical study input",
        "operational_authorization": {
            "usable_as_live_river_stage": False,
            "usable_as_tide_or_coastal_water": False,
            "usable_as_issue_time_model_feature": False,
            "reason": "Historical supporting information lacks a verified live station feed, operational datum validation, and prospective issue-time lineage.",
        },
        "sheets": sheets,
        "safety": "Monitoring only: inspection and provenance audit only; no model fitting, score, probability, prediction, forecast, alert, or life-safety decision.",
    }
    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
