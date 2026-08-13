from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\geo-ai-floating-predection")
OUT = ROOT / "outputs"
TEAM_CODE = "HACK-MM-001"
PROJECT_TITLE = "FloodGuard Myanmar: GeoAI Flood Susceptibility and Early Warning MVP"


def set_font(run, name="Calibri", size=11, bold=False, color="222222"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, size=11, bold=False, color="222222"):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_font(r)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(header)
        set_font(r, bold=True, color="0B2545")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(value)
            set_font(r, size=10)
    doc.add_paragraph()
    return table


def build_pdgs():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = add_para(doc, PROJECT_TITLE, size=20, bold=True, color="0B2545")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = add_para(doc, f"{TEAM_CODE}_PDGS_v01 | ASEAN GeoAI Fusion 2026 Virtual Hackathon", size=10, color="555555")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_para(
        doc,
        "Concise Overview",
        style="Heading 1",
        size=15,
        bold=True,
        color="1F4D78",
    )
    add_para(
        doc,
        "FloodGuard Myanmar is a GeoAI and IoT-informed flood intelligence MVP for Maubin Township, Ayeyarwady Region. "
        "It combines township terrain, waterways, land cover, rainfall history, historical flood evidence, weather forecasts, "
        "and future sensor ingestion to help local responders identify flood-prone areas and prioritize monitoring before a full operational warning system is available.",
    )

    add_para(doc, "Problem", style="Heading 1", size=15, bold=True, color="1F4D78")
    add_bullets(
        doc,
        [
            "Flood-prone communities often lack localized, timely, and easy-to-understand flood risk information.",
            "Satellite and reanalysis datasets help regional screening but are delayed, coarse, or historical; they cannot replace local observations.",
            "Decision makers need transparent map outputs that distinguish historical susceptibility, rainfall context, sensor thresholds, and operational limitations.",
        ],
    )

    add_para(doc, "Target Users and Beneficiaries", style="Heading 1", size=15, bold=True, color="1F4D78")
    add_bullets(
        doc,
        [
            "Township disaster management teams and local administrators.",
            "Community volunteers and village-level flood monitors.",
            "Residents in low-lying or waterway-adjacent settlements.",
            "Humanitarian, planning, and infrastructure teams that need spatial prioritization evidence.",
        ],
    )

    add_para(doc, "Datasets and Data Sources", style="Heading 1", size=15, bold=True, color="1F4D78")
    add_table(
        doc,
        ["Data Source", "Use in MVP", "Current Notes"],
        [
            ["MIMU/UNOSAT township boundary", "Clipping, map outline, analysis extent", "Public-use permission should be confirmed before public deployment."],
            ["Copernicus DEM GLO-30", "Elevation, relief, 500 m terrain cells", "5,549 township-clipped cells; DSM limitations in low-relief areas."],
            ["OpenStreetMap waterways", "River/canal proximity and map overlays", "282 river segments and 21 canal segments; ODbL attribution required."],
            ["ESA WorldCover 2021", "Land-cover shares per terrain cell", "10 m reference snapshot; temporal mismatch for older floods."],
            ["ERA5 rainfall history", "Daily rainfall history and rolling accumulations", "9,497 daily rows covering 2000-01-01 to 2025-12-31."],
            ["Global Flood Database / MODIS", "Historical flood frequency and event labels", "17 Maubin-overlap events; CC BY-NC 4.0 terms require care."],
            ["Open-Meteo forecast", "1-7 day rainfall outlook", "Forecast rainfall only; not a direct flood forecast."],
            ["ESP32/MQTT/HTTP sensor pipeline", "Future water-level and rainfall ingestion", "API-ready; physical deployment still pending."],
        ],
    )

    add_para(doc, "GeoAI Solution Approach", style="Heading 1", size=15, bold=True, color="1F4D78")
    add_bullets(
        doc,
        [
            "Prepare township-scale spatial features in PostGIS from DEM, waterways, land cover, rainfall, and historical flood layers.",
            "Generate an explainable terrain screening score using low relative elevation, waterway proximity, and local flatness.",
            "Train historical susceptibility and rainfall-aligned event models on grid/event data, then expose model probabilities and hindcast outputs through FastAPI.",
            "Render evidence layers in a Next.js dashboard with 2D Leaflet, MapLibre 3D terrain, and Cesium globe views.",
            "Use future IoT water-level observations and threshold rules for live warnings once physical stations are deployed and calibrated.",
        ],
    )

    add_para(doc, "Expected Outcomes, Impact, and Scalability", style="Heading 1", size=15, bold=True, color="1F4D78")
    add_bullets(
        doc,
        [
            "Short-term: identify high-priority cells and waterways for field verification, sensor placement, and community monitoring.",
            "Prototype output: dashboard-ready risk layers, rainfall context, provenance catalog, historical event selector, and API-backed map inspection.",
            "Impact value: clearer local planning evidence before flood season and better communication of uncertainty and data limitations.",
            "Scalability: repeat the pipeline for other townships by replacing boundary, DEM clip, waterways, land cover, rainfall, flood labels, and sensor station configuration.",
        ],
    )

    add_para(doc, "Limitations and Risk Controls", style="Heading 1", size=15, bold=True, color="1F4D78")
    add_bullets(
        doc,
        [
            "Current outputs are historical susceptibility and hindcast evidence, not a certified operational flood warning.",
            "Flood depth, arrival time, tide, levee, drainage capacity, and calibrated physical sensors are not yet implemented.",
            "Model false positives remain a risk; human review and official warning channels are required for operational use.",
            "All dataset licences and attributions must be verified before public or commercial deployment.",
        ],
    )

    add_para(doc, "Submission Metadata", style="Heading 1", size=15, bold=True, color="1F4D78")
    add_table(
        doc,
        ["Field", "Suggested Value"],
        [
            ["Project title", PROJECT_TITLE],
            ["Short summary", "GeoAI flood susceptibility and early-warning MVP for Maubin Township using PostGIS, historical flood evidence, rainfall data, 2D/3D map visualization, and future IoT sensor ingestion."],
            ["Repository URL", "Add your GitHub/GitLab repository URL before final submission."],
            ["AI usage declaration", "AI tools were used to assist drafting, coding support, and submission material preparation; the team remains responsible for validation and final content."],
            ["Dataset usage rights declaration", "Datasets are public/open or research-access sources with attribution and licence checks required before public deployment."],
            ["Originality/IP declaration", "Prototype source code and integration work are team-created unless otherwise noted; third-party datasets and libraries retain their original licences."],
            ["Compliance declaration", "The team should confirm compliance with ASEAN GeoAI Fusion 2026 terms before final submission."],
        ],
    )

    out = OUT / f"{TEAM_CODE}_PDGS_v01.docx"
    doc.save(out)
    return out


def build_text_assets():
    script = OUT / f"{TEAM_CODE}_PRES_v01_script.txt"
    script.write_text(
        """Title: HACK-MM-001_PRES_v01

5-minute recorded presentation script

0:00-0:25 Team and project
Hello, we are presenting FloodGuard Myanmar, a GeoAI flood susceptibility and early-warning MVP for Maubin Township in Ayeyarwady Region. The project focuses on turning geospatial evidence, rainfall information, historical flood labels, and future local sensor readings into clear map-based decision support.

0:25-1:05 Problem
Flood-prone communities need localized and understandable information before water becomes dangerous. Satellite and reanalysis data are useful, but they are delayed, coarse, or historical. Local teams therefore need a system that separates current conditions, historical susceptibility, forecast rainfall, and model limitations instead of presenting one black-box warning.

1:05-2:10 Data and GeoAI method
Our MVP builds a PostGIS spatial dataset for Maubin Township. It uses township boundaries, Copernicus DEM GLO-30 elevation, OpenStreetMap waterways, ESA WorldCover land cover, ERA5 rainfall history, Open-Meteo forecast rainfall, and Global Flood Database historical events. The first model layer is an explainable terrain screening score based on low relative elevation, proximity to waterways, and local flatness. The second layer adds historical flood susceptibility and rainfall-aligned event hindcast models for grid cells and historical flood events.

2:10-3:35 Prototype demonstration
In the dashboard, users can open the Maubin map, switch between elevation, susceptibility, land cover, waterways, historical flood, ML probability, and event hindcast layers, then click a cell to inspect the evidence behind the score. The same API data can be viewed in 2D Leaflet, MapLibre 3D terrain, and Cesium globe views. The backend exposes FastAPI endpoints for health, spatial assets, terrain screening, rainfall history, flood extents, flood ML outputs, sensor stations, observations, alerts, and live WebSocket updates.

3:35-4:25 Benefits and impact
The value is practical prioritization. Local responders can identify which cells and waterways need field verification, where sensors should be placed, and which communities may need closer monitoring. The system also makes provenance and limitations visible, which reduces the risk of treating an experimental model as an official warning.

4:25-5:00 Limitations and next steps
This MVP is not yet a certified operational forecast. It does not yet predict flood depth, arrival time, levee failure, drainage capacity, tide effects, or calibrated live sensor forecasts. The next steps are physical sensor deployment, external validation, false-alarm reduction, human review workflows, messaging channels, authentication, and expansion to other townships. The approach is scalable because the same data pipeline can be repeated with new boundaries, local observations, and updated model labels.
""",
        encoding="utf-8",
    )

    fields = OUT / f"{TEAM_CODE}_submission_form_fields_v01.txt"
    fields.write_text(
        f"""Suggested form fields for ASEAN GeoAI Fusion 2026

Project title:
{PROJECT_TITLE}

Short summary:
GeoAI flood susceptibility and early-warning MVP for Maubin Township using PostGIS, historical flood evidence, rainfall data, 2D/3D map visualization, and future IoT sensor ingestion.

Repository URL:
Add your repository URL here before final submission.

Recorded presentation title:
{TEAM_CODE}_PRES_v01

Recorded presentation link:
Record the 5-minute presentation using the provided script, upload it to an accessible link, and paste that link into the submission form.

AI Usage Declaration:
AI tools were used to assist with drafting, coding support, and preparation of submission materials. The team reviewed and remains responsible for the final technical claims, prototype behavior, and submitted content.

Dataset Usage Rights Declaration:
The project uses public or open geospatial, weather, satellite, and community-mapped datasets with attribution and licence constraints. The team should confirm public-use permission and licence compliance before final publication or deployment.

Originality and Intellectual Property Declaration:
The prototype integration, application code, and submission materials are original team work unless noted. Third-party datasets, map tiles, libraries, and services retain their original licences and attribution requirements.

Compliance with AGAIF Terms and Conditions:
The team should review and confirm all ASEAN GeoAI Fusion 2026 requirements, terms, deadlines, file naming, versioning, and accessibility requirements before selecting Submit Final.
""",
        encoding="utf-8",
    )
    return script, fields


def build_proto_zip():
    zip_path = OUT / f"{TEAM_CODE}_PROTO_v01.zip"
    excludes = {
        ".git",
        ".pnpm-store",
        "node_modules",
        ".next",
        ".venv",
        "__pycache__",
        ".pytest_cache",
        "outputs",
    }
    with ZipFile(zip_path, "w", ZIP_DEFLATED) as zf:
        for path in ROOT.rglob("*"):
            rel = path.relative_to(ROOT)
            if any(part in excludes for part in rel.parts):
                continue
            if path.is_file():
                zf.write(path, rel.as_posix())
    return zip_path


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    print(build_pdgs())
    print(build_text_assets())
    print(build_proto_zip())
